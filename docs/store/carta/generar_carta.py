"""Carta de Arta Producciones que autoriza a NEXARA a publicar la app con el nombre y el logo de ARTA.

Apple la pide por la guía 5.2.1 (contenido o marca de un tercero) y Google la acepta como prueba de
autorización. Sale en Word (para firmar) y en PDF, en español con traducción al inglés para el revisor.
Los datos del firmante se dejan en blanco: los llena y firma el representante legal de Arta.

Uso:  python docs/store/carta/generar_carta.py [carpeta-de-salida]
      (por omisión, Documentos\\ARTA-builds)
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[3]
LOGO = ROOT / "apps/web/public/brand/arta-logo-ink.png"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Documents" / "ARTA-builds"
NAME = "Carta-autorizacion-ARTA-NEXARA"

ORG = "Arta Producciones S.A. de C.V."
PUB = "NEW ENGINEERING EXPERTISE AND RESOURCE ADVANCEMENT S.A. DE C.V."
APP = "ARTA Producciones"
BUNDLE = "com.artaproducciones.ops"
TEAM = "AHNW9K8745"

ES = {
    "lugar": "Puebla, Puebla, a 4 de octubre de 2026",
    "para": "A quien corresponda — Apple Inc. (App Review) y Google LLC (Google Play)",
    "asunto": f"Asunto: autorización para publicar la aplicación «{APP}»",
    "cuerpo": [
        f"{ORG} («Arta Producciones»), titular del nombre comercial, la marca y el logotipo ARTA, autoriza a "
        f"{PUB} («NEXARA»), equipo de Apple Developer {TEAM}, a desarrollar, publicar, actualizar y mantener en "
        f"App Store y en Google Play la aplicación «{APP}» (identificador {BUNDLE}) usando el nombre, la marca y "
        "el logotipo de ARTA.",
        "NEXARA desarrolla la aplicación por encargo de Arta Producciones y la opera como su proveedor de "
        "tecnología. La aplicación es la herramienta de trabajo de nuestro equipo y de las empresas de producción "
        "de eventos que usan la plataforma ARTA: tareas, chat del equipo, aprobaciones, órdenes de compra y "
        "anticipos.",
        "Esta autorización está vigente mientras exista la relación de servicios entre ambas empresas. Arta "
        "Producciones puede revocarla en cualquier momento mediante aviso por escrito.",
        "Para cualquier aclaración: contacto@artaproducciones.com · https://artaproducciones.com",
    ],
    "atte": "Atentamente,",
    "firma": ["Nombre: ________________________________", "Cargo: Representante legal", ORG],
}
EN = {
    "lugar": "Puebla, Mexico, October 4, 2026",
    "para": "To whom it may concern — Apple Inc. (App Review) and Google LLC (Google Play)",
    "asunto": f"Re: authorization to publish the “{APP}” app",
    "cuerpo": [
        f"{ORG} (“Arta Producciones”), owner of the ARTA trade name, trademark and logo, authorizes {PUB} "
        f"(“NEXARA”), Apple Developer team {TEAM}, to develop, publish, update and maintain on the App Store and "
        f"Google Play the “{APP}” app (identifier {BUNDLE}) using the ARTA name, trademark and logo.",
        "NEXARA develops the app on behalf of Arta Producciones and operates it as our technology provider. The "
        "app is the work tool of our team and of the event production companies that use the ARTA platform: "
        "tasks, team chat, approvals, purchase orders and cash advances.",
        "This authorization remains in effect while the service relationship between both companies continues. "
        "Arta Producciones may revoke it at any time by written notice.",
        "Contact: contacto@artaproducciones.com · https://artaproducciones.com",
    ],
    "atte": "Sincerely,",
    "firma": ["Name: ________________________________", "Title: Legal representative", ORG],
    "nota": "English translation of the letter above.",
}


def word(path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.5))
    sec.top_margin = sec.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def page(t, nota=None):
        doc.add_picture(str(LOGO), width=Cm(4.2))
        if nota:
            p = doc.add_paragraph(nota)
            p.runs[0].italic = True
            p.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
        r = doc.add_paragraph().add_run(t["lugar"])
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph(t["para"])
        doc.add_paragraph().add_run(t["asunto"]).bold = True
        for parrafo in t["cuerpo"]:
            p = doc.add_paragraph(parrafo)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_paragraph(t["atte"])
        doc.add_paragraph("\n\n\n_________________________________________")
        for linea in t["firma"]:
            doc.add_paragraph(linea)

    page(ES)
    doc.add_page_break()
    page(EN, EN["nota"])
    doc.save(path)


def pdf(path: Path):
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=11, leading=15.5, textColor=HexColor("#111111"))
    just = ParagraphStyle("just", parent=base, alignment=TA_JUSTIFY, spaceAfter=9)
    right = ParagraphStyle("right", parent=base, alignment=2, spaceAfter=14)
    bold = ParagraphStyle("bold", parent=base, fontName="Helvetica-Bold", spaceBefore=6, spaceAfter=12)
    note = ParagraphStyle("note", parent=base, fontName="Helvetica-Oblique", fontSize=9, textColor=HexColor("#707070"))

    logo = Image(str(LOGO))
    ratio = logo.imageHeight / logo.imageWidth
    flow = []

    def page(t, nota=None):
        img = Image(str(LOGO), width=4.2 * cm, height=4.2 * cm * ratio)
        img.hAlign = "LEFT"
        flow.extend([img, Spacer(1, 12)])
        if nota:
            flow.append(Paragraph(nota, note))
        flow.extend([Paragraph(t["lugar"], right), Paragraph(t["para"], base), Paragraph(t["asunto"], bold)])
        flow.extend(Paragraph(p, just) for p in t["cuerpo"])
        flow.extend([Spacer(1, 6), Paragraph(t["atte"], base), Spacer(1, 52)])
        flow.append(Paragraph("_________________________________________", base))
        flow.extend(Paragraph(linea, base) for linea in t["firma"])

    page(ES)
    flow.append(PageBreak())
    page(EN, EN["nota"])
    SimpleDocTemplate(
        str(path), pagesize=LETTER, leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=f"Autorización — {APP}", author=ORG,
    ).build(flow)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    word(OUT / f"{NAME}.docx")
    pdf(OUT / f"{NAME}.pdf")
    print("listo:", OUT / f"{NAME}.docx", "y", OUT / f"{NAME}.pdf")
